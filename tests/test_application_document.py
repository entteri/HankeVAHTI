from io import BytesIO
from zipfile import ZipFile

from docx import Document
import pytest

from app.services.application_document import (
    ApplicationDocumentError, TEMPLATE_PATH, application_filename, export_application_document,
)
from app.services.application_template import APPLICATION_FIELDS, KNOWN_FIELDS
from app.services.funding_calls import get_funding_call
from tests.test_ai_summary import _selected_call


@pytest.fixture
def call(db_session):
    model = _selected_call(db_session)
    model.title = "Haku / ../ <testi>"
    model.call_identifier = "EURA-2026-123"
    db_session.commit()
    return get_funding_call(db_session, model.id)


def test_every_field_maps_to_correct_answer_and_template_unchanged(call):
    original = TEMPLATE_PATH.read_bytes()
    content = {"project_context": "EI-WORDIIN", "fields": {
        field.id: f"Vastaus {field.id}: ääkköset, ensimmäinen rivi.\nToinen rivi."
        for field in APPLICATION_FIELDS}, "additional_information": {"need": "EI-LISÄTIETOJA-WORDIIN"}}
    output = export_application_document(call, content)
    before, after = Document(BytesIO(original)), Document(BytesIO(output))
    assert len(after.tables) == 77 and len(after.paragraphs) == 202
    for field in APPLICATION_FIELDS:
        w = field.word
        assert after.tables[w.table].cell(w.row, w.column).paragraphs[w.paragraph].text == (
            content["fields"][field.id] + before.tables[w.table].cell(w.row, w.column).paragraphs[w.paragraph].text[len(w.placeholder):])
        assert after.tables[w.table].cell(w.label_row, w.label_column).paragraphs[w.label_paragraph].text == field.label
    assert after.tables[0].cell(1, 0).paragraphs[1].text == call.title
    assert after.tables[0].cell(1, 1).paragraphs[1].text == call.call_identifier
    assert TEMPLATE_PATH.read_bytes() == original
    # All unrelated cells (including financial/admin/date/legal fields) remain byte-identical XML.
    targets = {(f.word.table, f.word.row, f.word.column) for f in (*APPLICATION_FIELDS, *KNOWN_FIELDS)}
    for ti, table in enumerate(before.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if (ti, ri, ci) not in targets:
                    assert cell._tc.xml == after.tables[ti].cell(ri, ci)._tc.xml
    assert [p._p.xml for p in before.paragraphs] == [p._p.xml for p in after.paragraphs]
    assert before.sections[0]._sectPr.xml == after.sections[0]._sectPr.xml
    # Header/footer, styles, images and relationships remain original.
    with ZipFile(BytesIO(original)) as old, ZipFile(BytesIO(output)) as new:
        assert set(old.namelist()) == set(new.namelist())
        for name in old.namelist():
            if name != "word/document.xml":
                assert old.read(name) == new.read(name), name


def test_empty_fields_keep_placeholders_and_run_styles(call):
    output = export_application_document(call, {"fields": {"need": "", "objectives": "Täytetty tavoite"}})
    old, new = Document(TEMPLATE_PATH), Document(BytesIO(output))
    for field in APPLICATION_FIELDS:
        w = field.word
        a = old.tables[w.table].cell(w.row, w.column).paragraphs[w.paragraph]
        b = new.tables[w.table].cell(w.row, w.column).paragraphs[w.paragraph]
        if field.id != "objectives":
            assert a._p.xml == b._p.xml
        else:
            assert a.style.name == b.style.name
            assert len(a.runs) == len(b.runs)
            for ar, br in zip(a.runs, b.runs):
                assert (ar._r.rPr.xml if ar._r.rPr is not None else None) == (br._r.rPr.xml if br._r.rPr is not None else None)


@pytest.mark.parametrize("corruption", ["question", "placeholder", "table", "missing", "invalid"])
def test_wrong_template_is_safe_error_even_for_empty_field(call, tmp_path, corruption):
    path = tmp_path / "template.docx"
    document = Document(TEMPLATE_PATH)
    if corruption == "missing":
        pass
    elif corruption == "invalid":
        path.write_text("not docx")
    else:
        if corruption == "question":
            document.tables[20].cell(0, 0).paragraphs[0].text = "Väärä kysymys"
        elif corruption == "placeholder":
            document.tables[20].cell(1, 0).paragraphs[0].text = "Jo täytetty"
        else:
            element = document.tables[20]._tbl
            element.getparent().remove(element)
        document.save(path)
    before = path.read_bytes() if path.exists() else None
    with pytest.raises(ApplicationDocumentError):
        export_application_document(call, {}, template_path=path)
    assert (path.read_bytes() if path.exists() else None) == before


def test_split_placeholder_preserves_question_and_formatting(call, tmp_path):
    document = Document(TEMPLATE_PATH)
    answer = document.tables[20].cell(1, 0).paragraphs[0]
    answer.clear()
    answer.add_run("Kirjoita ").bold = True
    answer.add_run("vastaus tähän\n\n").italic = True
    path = tmp_path / "split.docx"
    document.save(path)
    output = Document(BytesIO(export_application_document(call, {"fields": {"objectives": "Tavoite"}}, template_path=path)))
    result = output.tables[20].cell(1, 0).paragraphs[0]
    assert result.text == "Tavoite\n\n"
    assert result.runs[0].bold and result.runs[1].italic


def test_filename_only_uses_numeric_id(call):
    assert application_filename(call.id) == f"Hankehakemus_{call.id}.docx"
    with pytest.raises(ValueError):
        application_filename("../../bad")
