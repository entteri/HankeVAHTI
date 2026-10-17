import asyncio
import json
from io import BytesIO

from docx import Document
from nicegui import events, ui
import pytest
from sqlalchemy.orm import sessionmaker

from app.services.eura_application_drafts import generate_application_field, get_application_draft
from app.services.application_text import ApplicationTextError, ApplicationTextResult
from tests.test_ai_summary import _selected_call
from tests.test_funding_calls import _test_client
from tests.test_relevance_ui import _click, _page_client


@pytest.fixture
def setup(db_session, monkeypatch):
    call = _selected_call(db_session)
    monkeypatch.setattr("app.ui.application_draft.SessionLocal", sessionmaker(bind=db_session.get_bind()))
    # Any unmocked generation fails the test rather than reaching a provider.
    monkeypatch.setattr("app.ui.application_draft.generate_application_field", lambda *a, **k: pytest.fail("Unexpected AI request"))
    return call.id


def editor(page, key):
    return next(e for e in page.elements.values() if isinstance(e, ui.textarea) and e.props.get("data-field-id") == key)


def edit(page, key, value):
    with page:
        editor(page, key).set_value(value)


async def generate(page, key):
    button = next(e for e in page.elements.values() if isinstance(e, ui.button) and e.props.get("data-field-id") == key)
    assert button.enabled
    with page:
        listener = next(e for e in button._event_listeners.values() if e.type == "click")
        before = asyncio.all_tasks()
        events.handle_event(listener.handler, events.GenericEventArguments(sender=button, client=page, args=None))
        pending = asyncio.all_tasks() - before
        if pending:
            await asyncio.wait_for(asyncio.gather(*pending), timeout=5)


def test_details_link_and_old_formatter_remain(db_session, monkeypatch, setup):
    navigated = []
    monkeypatch.setattr(ui.navigate, "to", lambda path: navigated.append(path))
    with _test_client(db_session, monkeypatch) as client:
        page = _page_client(client.get("/hankkeet"))
        assert any(isinstance(e, ui.button) and e.text == "Muotoile hakemustekstiä" for e in page.elements.values())
        client.portal.call(_click, page, "Luo hankehakemus")
        assert navigated == [f"/hankkeet/{setup}/hakemus"]
        assert get_application_draft(db_session, setup) is None


def test_edit_save_reload_and_download_saved_word(db_session, monkeypatch, setup):
    downloads = []
    monkeypatch.setattr(ui, "download", lambda data, **kwargs: downloads.append((data, kwargs)))
    with _test_client(db_session, monkeypatch) as client:
        page = _page_client(client.get(f"/hankkeet/{setup}/hakemus"))
        assert get_application_draft(db_session, setup) is None
        def context():
            with page:
                next(e for e in page.elements.values() if isinstance(e, ui.textarea) and e.props.get("label") == "Omat ideat ja tarkennukset (vapaaehtoinen)").set_value("Hankeidea")
        client.portal.call(context)
        client.portal.call(edit, page, "need", "Käyttäjän tarve")
        client.portal.call(edit, page, "risk_1_risk", "Alustava riski")
        client.portal.call(_click, page, "Poista riski 1")
        assert editor(page, "risk_1_risk").value == ""
        client.portal.call(_click, page, "Lataa Word-luonnos")
        assert not downloads
        client.portal.call(_click, page, "Tallenna luonnos")
        saved = get_application_draft(db_session, setup)
        assert saved.content_json["project_context"] == "Hankeidea"
        assert saved.content_json["fields"]["need"] == "Käyttäjän tarve"
        reloaded = _page_client(client.get(f"/hankkeet/{setup}/hakemus"))
        assert editor(reloaded, "need").value == "Käyttäjän tarve"
        client.portal.call(_click, reloaded, "Lataa Word-luonnos")
        assert len(downloads) == 1
        data, kwargs = downloads[0]
        assert kwargs["filename"] == f"Hankehakemus_{setup}.docx"
        doc = Document(BytesIO(data))
        assert doc.tables[18].cell(1, 0).text.startswith("Käyttäjän tarve")


@pytest.mark.parametrize("outcome", ["success", "additional_only", "error", "unexpected"])
def test_field_ai_only_updates_target_and_keeps_editable(db_session, monkeypatch, setup, outcome):
    calls = []
    def request(call, key, content):
        calls.append((key, content))
        if outcome == "error":
            raise ApplicationTextError("Palvelun testivirhe")
        if outcome == "unexpected":
            raise RuntimeError("Do not expose provider internals")
        return ApplicationTextResult("AI:n tavoite" if outcome == "success" else "", "Tarkista lähtötieto", None)
    monkeypatch.setattr("app.ui.application_draft.generate_application_field", request)
    with _test_client(db_session, monkeypatch) as client:
        page = _page_client(client.get(f"/hankkeet/{setup}/hakemus"))
        client.portal.call(edit, page, "objectives", "Oma tavoite")
        client.portal.call(edit, page, "need", "Muu teksti")
        client.portal.call(generate, page, "objectives")
        assert len(calls) == 1 and calls[0][0] == "objectives"
        assert calls[0][1]["fields"]["objectives"] == "Oma tavoite"
        assert editor(page, "need").value == "Muu teksti"
        assert editor(page, "objectives").value == ("AI:n tavoite" if outcome == "success" else "Oma tavoite")
        assert editor(page, "objectives").enabled
        assert get_application_draft(db_session, setup) is None  # generation is not autosave
        if outcome in ("success", "additional_only"):
            assert "Tarvittavat lisätiedot: Tarkista lähtötieto" in [e.text for e in page.elements.values() if isinstance(e, ui.label)]
        client.portal.call(edit, page, "objectives", "Muokattu vielä")
        client.portal.call(_click, page, "Tallenna luonnos")
        assert get_application_draft(db_session, setup).content_json["fields"]["objectives"] == "Muokattu vielä"


def test_missing_call_page(db_session, monkeypatch, setup):
    with _test_client(db_session, monkeypatch) as client:
        page = _page_client(client.get("/hankkeet/99999/hakemus"))
        assert "Rahoitushakua ei löydy." in [e.text for e in page.elements.values() if isinstance(e, ui.label)]


def test_confirmation_persists_and_edit_revokes_it(db_session, monkeypatch, setup):
    with _test_client(db_session, monkeypatch) as client:
        page = _page_client(client.get(f"/hankkeet/{setup}/hakemus"))
        client.portal.call(edit, page, "project_name", "Oma hanke")
        check = next(e for e in page.elements.values() if isinstance(e, ui.checkbox))
        def confirm():
            with page:
                check.set_value(True)
        client.portal.call(confirm)
        client.portal.call(_click, page, "Tallenna luonnos")
        assert get_application_draft(db_session, setup).content_json["confirmed_fields"] == ["project_name"]
        client.portal.call(edit, page, "project_name", "Muuttunut hanke")
        assert check.value is False
        client.portal.call(_click, page, "Tallenna luonnos")
        db_session.expire_all()
        assert get_application_draft(db_session, setup).content_json["confirmed_fields"] == []


@pytest.mark.parametrize("field_id", ["actions", "public_summary"])
def test_ai_works_without_ideas_and_does_not_auto_confirm(db_session, monkeypatch, setup, field_id):
    monkeypatch.setattr("app.core.config.AI_PROVIDER", "gemini")
    monkeypatch.setattr("app.ui.application_draft.generate_application_field", generate_application_field)
    requests = []
    def gemini(instructions, message, **kwargs):
        data = json.loads(message)
        requests.append(data)
        assert data["project_context"] == ""
        assert data["operation"] == "suggest"
        assert data["field_id"] == field_id
        assert data["funding_call"]["description"] == "Hanke kehittää henkilöstön digitaitoja."
        assert "Käyttäjän omat ideat ja tarkennukset (project_context) ovat vapaaehtoisia." in instructions
        if field_id == "public_summary":
            assert "## Hankkeen julkinen tiivistelmä (public_summary)" in instructions
            assert "muodosta yksi yhtenäinen alustava hanke-ehdotus" in instructions
            assert "vaatimus valmiista hanketiedoista ei estä" in instructions
        return json.dumps({"text": "Hankkeessa voitaisiin järjestää työpajoja.", "additional_information": "Vahvista toteutustapa."})
    monkeypatch.setattr("app.services.application_text.generate_gemini_text", gemini)
    with _test_client(db_session, monkeypatch) as client:
        page = _page_client(client.get(f"/hankkeet/{setup}/hakemus"))
        assert any(isinstance(e, ui.textarea) and e.props.get("label") == "Omat ideat ja tarkennukset (vapaaehtoinen)" and e.value == "" for e in page.elements.values())
        client.portal.call(generate, page, field_id)
        assert len(requests) == 1
        assert editor(page, field_id).value == "Hankkeessa voitaisiin järjestää työpajoja."
        assert editor(page, field_id).enabled
        client.portal.call(_click, page, "Tallenna luonnos")
        content = get_application_draft(db_session, setup).content_json
        assert content["project_context"] == ""
        assert content["confirmed_fields"] == []
        assert content["additional_information"][field_id] == "Vahvista toteutustapa."
