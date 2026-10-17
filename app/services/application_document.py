"""Fill only verified answer locations in an in-memory copy of the original DOCX."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from app.core.config import PROJECT_ROOT
from app.services.eura_application_drafts import EuraApplicationDraftError, normalise_content
from app.services.application_template import APPLICATION_FIELDS, KNOWN_FIELDS
from app.services.funding_calls import FundingCallView

TEMPLATE_PATH = PROJECT_ROOT / "templates" / "hankehakemus_tyhja_pohja.docx"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ApplicationDocumentError(Exception):
    """A Finnish export error safe to display without document or API contents."""


def _targets(document):
    if len(document.tables) != 77 or len(document.paragraphs) != 202:
        raise ApplicationDocumentError("Word-pohjan rakenne on muuttunut. Tarkista pohja ja kenttäkohdistukset.")
    targets = {}
    for field in (*KNOWN_FIELDS, *APPLICATION_FIELDS):
        w = field.word
        table = document.tables[w.table]
        if (len(table.rows), len(table.columns)) != w.shape:
            raise ApplicationDocumentError(f"Word-pohjan taulukko ei vastaa kenttää: {field.label}")
        label = table.cell(w.label_row, w.label_column).paragraphs[w.label_paragraph].text
        paragraph = table.cell(w.row, w.column).paragraphs[w.paragraph]
        if label != field.label or paragraph.text.rstrip("\n") != w.placeholder:
            raise ApplicationDocumentError(f"Word-pohjan kysymys tai vastausalue on muuttunut: {field.label}")
        targets[field.id] = paragraph
    return targets


def _replace_placeholder(paragraph, placeholder: str, value: str) -> None:
    """Retain paragraph properties, run formatting and trailing template line breaks."""
    offset = 0
    inserted = False
    for run in paragraph.runs:
        original = run.text
        end = offset + len(original)
        if offset < len(placeholder) and end > 0:
            suffix = original[max(0, len(placeholder) - offset):]
            run.text = (value if not inserted else "") + suffix
            inserted = True
        offset = end
    if not inserted:
        raise ApplicationDocumentError("Word-pohjan vastausalueen tekstiä ei voitu kohdistaa.")


def export_application_document(
    call: FundingCallView, content: dict, *, template_path: Path | None = None,
) -> bytes:
    try:
        validated = normalise_content(content)
        # Never save to the template path; output exists only in memory.
        original = (template_path or TEMPLATE_PATH).read_bytes()
        document = Document(BytesIO(original))
        targets = _targets(document)  # validate ALL mappings before writing any answer
        values = {**validated["fields"], "call_title": call.title,
                  "call_identifier": call.call_identifier or call.source_id}
        for field in (*KNOWN_FIELDS, *APPLICATION_FIELDS):
            value = values.get(field.id, "")
            if value and value.strip():
                _replace_placeholder(targets[field.id], field.word.placeholder, value)
        modified = BytesIO()
        document.save(modified)
        # python-docx may normalise/drop unrelated package parts. Only the document
        # body was edited, so keep every other ZIP member exactly as supplied.
        output = BytesIO()
        with ZipFile(BytesIO(original)) as source, ZipFile(modified) as filled, ZipFile(output, "w") as target:
            for member in source.infolist():
                target.writestr(member, filled.read(member.filename) if member.filename == "word/document.xml" else source.read(member.filename))
        return output.getvalue()
    except ApplicationDocumentError:
        raise
    except EuraApplicationDraftError as exc:
        raise ApplicationDocumentError(str(exc)) from exc
    except Exception as exc:
        raise ApplicationDocumentError("Word-luonnoksen luonti epäonnistui. Tarkista alkuperäinen Word-pohja ja kenttien sisältö.") from exc


def application_filename(call_id: int) -> str:
    # Only the internal numeric id is used, never a user/provider supplied path or title.
    return f"Hankehakemus_{int(call_id)}.docx"
